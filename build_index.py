#!/usr/bin/env python3
"""
PART A - Knowledge preparation.  Run once, and again whenever documents change:

    python build_index.py

  knowledge/docs/*.pdf|docx|txt|md
        -> 1. read text (PDF pages, Word paragraphs, OCR for scanned pages if Tesseract is installed)
        -> 2. clean (remove repeating headers/footers & page numbers, fix hyphenation, whitespace)
        -> 3. chunk (~150-250 words at paragraph/heading boundaries, ~40-word overlap)
        -> 4. embed each chunk (bge-small-en-v1.5 -> 384 numbers per chunk)
        -> 5. keyword tokens for BM25
        -> 6. save to knowledge/index/  (plain JSON + one .npy file; no database server needed)
  knowledge/pinned.json -> 7. embed the pinned questions too
"""
import json
import logging
import re
import sys
import time
from collections import Counter
from pathlib import Path

import numpy as np

import config
from core.embedder import get_embedder
from core.textutil import split_sentences, tokenize

logging.basicConfig(level=logging.INFO, format="%(message)s")
log = logging.getLogger("build_index")


# ============================================================ 1. READ
def read_pdf(path):
    import fitz  # pymupdf
    pages = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc):
            text = page.get_text("text")
            if len(text.strip()) < 20:  # probably a scanned image page
                text = ocr_page(page) or text
            pages.append((i + 1, text))
    return pages


def ocr_page(page):
    try:
        import pytesseract
        from PIL import Image
        pix = page.get_pixmap(dpi=300)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        log.info("    OCR page %s", page.number + 1)
        return pytesseract.image_to_string(img, lang="eng")
    except Exception:  # noqa: BLE001
        return ""


def read_docx(path):
    import docx
    d = docx.Document(path)
    lines = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            lines.append("")
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        lines.append(("# " + t) if style.startswith("heading") or style == "title" else t)
    for table in d.tables:  # tables -> "cell | cell" lines
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
        lines.append("")
    return [(None, "\n".join(lines))]


def read_text(path):
    return [(None, Path(path).read_text(encoding="utf-8", errors="ignore"))]


READERS = {".pdf": read_pdf, ".docx": read_docx, ".txt": read_text, ".md": read_text}


# ============================================================ 2. CLEAN
def strip_headers_footers(pages):
    """Lines that repeat on more than half of the pages (headers/footers) are removed."""
    if len(pages) < 3:
        return pages
    counts = Counter()
    for _, text in pages:
        lines = {l.strip() for l in text.splitlines() if l.strip()}
        counts.update(lines)
    repeated = {l for l, c in counts.items() if c > len(pages) / 2 and len(l) < 120}
    return [(p, "\n".join(l for l in t.splitlines() if l.strip() not in repeated)) for p, t in pages]


def clean(text):
    text = text.replace("\r", "")
    text = re.sub(r"<!--.*?-->", " ", text, flags=re.S)               # editor notes in .md files
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)                      # hyphen-\nation -> hyphenation
    text = re.sub(r"^\s*(page\s*)?\d+(\s*(of|/)\s*\d+)?\s*$", "", text, flags=re.M | re.I)  # page numbers
    text = re.sub(r"[ \t ]+", " ", text)
    # join lines inside a paragraph (PDFs break lines mid-sentence), keep blank lines & headings
    out, buf = [], []
    for line in text.split("\n"):
        s = line.strip()
        if not s:
            if buf:
                out.append(" ".join(buf)); buf = []
            out.append("")
        elif s.startswith("#") or s.startswith("- ") or s.startswith("• "):
            if buf:
                out.append(" ".join(buf)); buf = []
            out.append(s)
        else:
            buf.append(s)
    if buf:
        out.append(" ".join(buf))
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out)).strip()


# ============================================================ 3. CHUNK
def paragraphs_with_headings(text):
    """Yield (heading_path, paragraph). Bullet lines are grouped into one paragraph."""
    title, section = None, None
    for block in re.split(r"\n\s*\n", text):
        block = block.strip()
        if not block:
            continue
        lines = block.split("\n")
        rest = []
        for line in lines:
            m = re.match(r"^(#+)\s*(.+)$", line.strip())
            if m:
                if rest:
                    yield (title, section), " ".join(rest); rest = []
                level, name = len(m.group(1)), m.group(2).strip()
                if level == 1 and title is None:
                    title = name
                else:
                    section = name
            else:
                rest.append(re.sub(r"^[-•]\s*", "", line.strip()))
        if rest:
            yield (title, section), " ".join(rest)


def chunk_document(text, file_name, page):
    """Group paragraphs into ~150-250-word chunks.
    - a new heading starts a new chunk if the current one is already big enough;
      otherwise the short section is merged in, with its heading kept inline ("Birds of the forest: ...")
    - when a chunk is cut in the middle of a section, the next one starts with the last ~40 words (overlap)
      so a fact split across the cut is still found whole in one of them."""
    chunks = []
    cur_words, chunk_head, last_head = [], None, None

    def flush():
        if cur_words:
            chunks.append({"head": chunk_head, "body": " ".join(cur_words), "file": file_name, "page": page})

    for head, para in paragraphs_with_headings(text):
        words = para.split()
        pieces = [words] if len(words) <= config.CHUNK_MAX_WORDS else _split_long(para)
        for piece in pieces:
            new_section = head != last_head
            too_big = len(cur_words) + len(piece) > config.CHUNK_MAX_WORDS
            big_enough = len(cur_words) >= config.CHUNK_MIN_WORDS
            if cur_words and (too_big or (new_section and big_enough)):
                overlap = cur_words[-config.CHUNK_OVERLAP_WORDS:] if not new_section else []
                flush()
                cur_words, chunk_head = list(overlap), head
            elif not cur_words:
                chunk_head = head
            elif new_section and head[1]:
                piece = [f"{head[1]}:"] + piece
            last_head = head
            cur_words.extend(piece)
    flush()
    # a very short last piece is merged into the previous chunk instead of standing alone
    if len(chunks) >= 2 and len(chunks[-1]["body"].split()) < 60:
        tail = chunks.pop()
        sec = tail["head"][1]
        chunks[-1]["body"] += " " + (f"{sec}: " if sec and sec != chunks[-1]["head"][1] else "") + tail["body"]
    return chunks


def _split_long(para):
    out, cur = [], []
    for s in split_sentences(para):
        w = s.split()
        if cur and len(cur) + len(w) > config.CHUNK_MAX_WORDS:
            out.append(cur); cur = []
        cur.extend(w)
    if cur:
        out.append(cur)
    return out


# ============================================================ MAIN
def main():
    t0 = time.time()
    config.INDEX_DIR.mkdir(parents=True, exist_ok=True)
    files = sorted(p for p in config.DOCS_DIR.rglob("*") if p.suffix.lower() in READERS and not p.name.startswith("."))
    if not files:
        log.error("No documents found in %s", config.DOCS_DIR)
        sys.exit(1)

    chunks = []
    for f in files:
        log.info("Reading %s", f.name)
        try:
            pages = READERS[f.suffix.lower()](f)
        except Exception as e:  # noqa: BLE001
            log.error("  ! could not read %s: %s", f.name, e)
            continue
        pages = strip_headers_footers(pages)
        doc_chunks = []
        for page_no, raw in pages:
            text = clean(raw)
            if text:
                doc_chunks.extend(chunk_document(text, f.name, page_no))
        log.info("  -> %d chunks", len(doc_chunks))
        chunks.extend(doc_chunks)

    for i, c in enumerate(chunks):
        title, section = c.pop("head") or (None, None)
        heading = " - ".join(x for x in (title, section) if x)
        c["id"] = i
        c["section"] = heading
        c["text"] = c.pop("body")
        # the heading is added in front for searching: "Satkosia Tiger Reserve - Animals of Satkosia: ..."
        c["search_text"] = (heading + ": " if heading else "") + c["text"]
        c["tokens"] = tokenize(c["search_text"])

    # pinned answers
    pinned = []
    if config.PINNED_FILE.exists():
        data = json.loads(config.PINNED_FILE.read_text(encoding="utf-8"))
        for item in data.get("items", []):
            for q in item.get("questions", []):
                pinned.append({"id": item["id"], "question": q, "answer": item["answer"], "tokens": tokenize(q)})

    # embeddings
    emb = get_embedder(allow_download=True)
    meta = {"built_at": time.strftime("%Y-%m-%d %H:%M:%S"), "n_chunks": len(chunks), "n_pinned": len(pinned),
            "files": [f.name for f in files], "embed_model": None}
    emb_file = config.INDEX_DIR / "embeddings.npy"
    pin_emb_file = config.INDEX_DIR / "pinned_embeddings.npy"
    if emb is not None:
        log.info("Embedding %d chunks with %s ...", len(chunks), emb.name)
        np.save(emb_file, emb.embed_docs([c["search_text"] for c in chunks]))
        if pinned:
            np.save(pin_emb_file, emb.embed_docs([p["question"] for p in pinned]))
        elif pin_emb_file.exists():
            pin_emb_file.unlink()
        meta["embed_model"] = emb.name
    else:
        log.warning("No embedding model available - index will be keyword-only (BM25).")
        for p in (emb_file, pin_emb_file):
            if p.exists():
                p.unlink()

    (config.INDEX_DIR / "chunks.json").write_text(json.dumps(chunks, ensure_ascii=False, indent=1), encoding="utf-8")
    (config.INDEX_DIR / "pinned_index.json").write_text(json.dumps(pinned, ensure_ascii=False, indent=1), encoding="utf-8")
    (config.INDEX_DIR / "meta.json").write_text(json.dumps(meta, indent=1), encoding="utf-8")
    words = [len(c["text"].split()) for c in chunks]
    log.info("Done: %d chunks (%d-%d words, avg %d), %d pinned questions, %.1fs -> %s",
             len(chunks), min(words), max(words), sum(words) // len(words), len(pinned), time.time() - t0, config.INDEX_DIR)


if __name__ == "__main__":
    main()
